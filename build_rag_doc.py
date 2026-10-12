# -*- coding: utf-8 -*-
"""把 RAG 模块开发流程的 Markdown 渲染为 Word 文档，并绘制数据流示意图。"""
import math
import re
import textwrap
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(r"E:\Code\AIProgect\travel assistant")
MD = ROOT / "RAG模块开发流程_v1.1.md"
OUT = ROOT / "RAG模块开发流程_v1.1.docx"
ASSETS = ROOT / "rag_assets"
ASSETS.mkdir(exist_ok=True)

FONT_REG = r"C:\Windows\Fonts\msyh.ttc"
FONT_BOLD = r"C:\Windows\Fonts\msyhbd.ttc"
BT = chr(96)
FENCE = BT * 3


def _font(size, bold=False):
    return ImageFont.truetype(FONT_BOLD if bold else FONT_REG, size)


def flow_png():
    W, H = 2240, 1140
    im = Image.new("RGB", (W, H), "white")
    d = ImageDraw.Draw(im)

    def box(x1, y1, x2, y2, title, sub="", fill="#EAF2F8", edge="#2F5597"):
        d.rounded_rectangle((x1, y1, x2, y2), radius=16, fill=fill, outline=edge, width=3)
        cx = (x1 + x2) // 2
        bb = d.textbbox((0, 0), title, font=_font(23, True))
        d.text((cx - (bb[2] - bb[0]) / 2, y1 + 14), title, font=_font(23, True), fill="#17202A")
        if sub:
            for k, ln in enumerate(textwrap.wrap(sub, width=max(8, int((x2 - x1) / 17)))[:3]):
                bb = d.textbbox((0, 0), ln, font=_font(16))
                d.text((cx - (bb[2] - bb[0]) / 2, y1 + 54 + k * 23), ln, font=_font(16), fill="#34495E")

    def arrow(a, b, color="#566573"):
        d.line([a, b], fill=color, width=4)
        ang = math.atan2(b[1] - a[1], b[0] - a[0])
        L = 15
        d.polygon([b,
                   (b[0] - L * math.cos(ang - .55), b[1] - L * math.sin(ang - .55)),
                   (b[0] - L * math.cos(ang + .55), b[1] - L * math.sin(ang + .55))], fill=color)

    def elbow(points, color="#566573"):
        for a, b in zip(points[:-2], points[1:-1]):
            d.line([a, b], fill=color, width=4)
        arrow(points[-2], points[-1], color)

    def dashed(x1, y1, x2, y2, label, color="#7D3C98"):
        for x in range(x1, x2, 18):
            d.line([(x, y1), (min(x + 10, x2), y1)], fill=color, width=2)
            d.line([(x, y2), (min(x + 10, x2), y2)], fill=color, width=2)
        for y in range(y1, y2, 18):
            d.line([(x1, y), (x1, min(y + 10, y2))], fill=color, width=2)
            d.line([(x2, y), (x2, min(y + 10, y2))], fill=color, width=2)
        d.text((x1 + 16, y1 + 8), label, font=_font(19, True), fill=color)

    dashed(20, 30, 2220, 620, "离线 · 索引构建（数据快照更新时执行）")
    row1 = [
        (70, "数据源", "Wikivoyage / Wikidata / GeoNames 许可与时间登记"),
        (600, "ingest 采集", "原始快照只追加 manifest.csv 记录来源"),
        (1130, "clean + align", "去噪、失效链接 地名与别名对齐"),
        (1660, "split 切分", "350-600 字 重叠 60-100 字"),
    ]
    for x, t, s in row1:
        box(x, 120, x + 470, 290, t, s, fill="#EBF5FB", edge="#2874A6")
    for i in range(3):
        arrow((row1[i][0] + 470, 205), (row1[i + 1][0], 205))

    row2 = [
        (320, "embed 向量化", "BGE-M3 API，内容哈希缓存 接口失败时本地模型兜底"),
        (940, "index 索引", "Chroma travel_kb_v1 + BM25 词法索引"),
        (1560, "IndexManifest", "索引版本绑定 数据快照与 git commit"),
    ]
    for x, t, s in row2:
        box(x, 400, x + 540, 570, t, s, fill="#E8F8F5", edge="#148F77")
    for i in range(2):
        arrow((row2[i][0] + 540, 485), (row2[i + 1][0], 485))
    elbow([(1895, 290), (1895, 345), (590, 345), (590, 400)])

    dashed(20, 660, 2220, 950, "在线 · 每次问答（检索预算 300 ms）")
    row3 = [
        (60, "query + 过滤下推", "城市 / 类别 / 时间"),
        (500, "向量 Top10 ∥ BM25 Top10", "两路并行召回"),
        (940, "RRF 融合", "k=60，同分优先向量"),
        (1380, "Top3 证据装配", "citations 绑定 source_id"),
        (1820, "Agent compose", "逐句引用校验 无证据不回答"),
    ]
    for x, t, s in row3:
        box(x, 730, x + 380, 900, t, s, fill="#FDEBD0", edge="#B9770E")
    for i in range(4):
        arrow((row3[i][0] + 380, 815), (row3[i + 1][0], 815))
    elbow([(1830, 570), (1830, 625), (250, 625), (250, 730)])

    d.text((60, 990), "指标：Recall@3 ≥95%　MRR ≥0.90　依据覆盖率 100%　检索阶段 P95 ≤300 ms",
           font=_font(24, True), fill="#1F4E78")
    d.text((60, 1040), "版本切换：先建 travel_kb_v2 并评测达标，再切换 RAG_COLLECTION；旧集合保留 7 天用于回滚。",
           font=_font(20), fill="#555555")
    p = ASSETS / "rag_flow.png"
    im.save(p)
    return p


INLINE = re.compile(r"(\*\*.+?\*\*|" + BT + r"[^" + BT + r"]+" + BT + r")")


def add_runs(p, text, size=None, bold=False):
    for part in INLINE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith(BT) and part.endswith(BT):
            r = p.add_run(part[1:-1])
            r.font.name = "Consolas"
            r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        else:
            r = p.add_run(part)
        if bold:
            r.bold = True
        if size:
            r.font.size = Pt(size)


def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd"); tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def shade_paragraph(paragraph, fill):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    pPr.append(shd)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in("w:tcMar")
    if tcMar is None:
        tcMar = OxmlElement("w:tcMar"); tcPr.append(tcMar)
    for m, v in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
        node = tcMar.find(qn("w:" + m))
        if node is None:
            node = OxmlElement("w:" + m); tcMar.append(node)
        node.set(qn("w:w"), str(v)); node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader"); el.set(qn("w:val"), "true"); trPr.append(el)


def add_table(doc, headers, rows):
    n = len(headers)
    t = doc.add_table(rows=1, cols=n)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.style = "Table Grid"
    t.autofit = False
    lens = [max([len(str(headers[i]))] + [len(str(r[i])) for r in rows if i < len(r)]) for i in range(n)]
    weights = [max(1.4, x ** 0.62) for x in lens]
    total = 16.4
    widths = [total * w / sum(weights) for w in weights]
    hdr = t.rows[0]
    set_repeat_table_header(hdr)
    for i, h in enumerate(headers):
        c = hdr.cells[i]
        c.text = ""
        p = c.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_runs(p, str(h), size=8.5, bold=True)
        for r in p.runs:
            r.font.color.rgb = RGBColor(255, 255, 255)
        set_cell_shading(c, "1F4E78")
        c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_margins(c)
        c.width = Cm(widths[i])
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        for i in range(n):
            val = str(row[i]) if i < len(row) else ""
            c = cells[i]
            c.text = ""
            p = c.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            add_runs(p, val, size=8.5)
            c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(c)
            if ri % 2:
                set_cell_shading(c, "F3F6F9")
            c.width = Cm(widths[i])
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


def split_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]


SEP = re.compile(r"^\|[\s:\-|]+\|$")


def parse_blocks(text):
    lines = text.split("\n")
    blocks = []
    i = 0
    while i < len(lines):
        s = lines[i].rstrip()
        if s.startswith(FENCE):
            i += 1
            buf = []
            while i < len(lines) and not lines[i].startswith(FENCE):
                buf.append(lines[i]); i += 1
            i += 1
            blocks.append(("code", buf)); continue
        if s.startswith("|") and i + 1 < len(lines) and SEP.match(lines[i + 1].strip()):
            headers = split_row(s)
            i += 2
            rows = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                rows.append(split_row(lines[i])); i += 1
            blocks.append(("table", (headers, rows))); continue
        m = re.match(r"^(#{1,4})\s+(.*)$", s)
        if m:
            blocks.append(("h%d" % len(m.group(1)), m.group(2))); i += 1; continue
        if not s.strip():
            i += 1; continue
        m = re.match(r"^-\s+\[ \]\s+(.*)$", s)
        if m:
            blocks.append(("check", m.group(1))); i += 1; continue
        if s.startswith("- "):
            blocks.append(("bullet", s[2:])); i += 1; continue
        m = re.match(r"^(\d+)\.\s+(.*)$", s)
        if m:
            blocks.append(("num", (m.group(1), m.group(2)))); i += 1; continue
        blocks.append(("para", s)); i += 1
    return blocks


def build():
    flow = flow_png()
    md = MD.read_text(encoding="utf-8")
    blocks = parse_blocks(md)

    doc = Document()
    sec = doc.sections[0]
    sec.top_margin = Cm(2.2); sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2.2); sec.right_margin = Cm(2.2)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(6)
    for name, size in [("Title", 22), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11)]:
        s = styles[name]
        s.font.name = "Microsoft YaHei"
        s._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        s.font.size = Pt(size); s.font.color.rgb = RGBColor(0, 0, 0); s.font.bold = True
        s.paragraph_format.space_before = Pt(12); s.paragraph_format.space_after = Pt(6)
        s.paragraph_format.keep_with_next = True
    styles["Title"].paragraph_format.space_before = Pt(0)
    tp = styles["Title"]._element.get_or_add_pPr()
    tb = tp.find(qn("w:pBdr"))
    if tb is not None:
        tp.remove(tb)

    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("旅游规划助手 RAG 模块详细开发流程  v1.1").font.size = Pt(8)

    title_done = False
    for kind, payload in blocks:
        if kind == "h1":
            if not title_done:
                p = doc.add_paragraph(style="Title")
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.add_run(payload)
                title_done = True
                continue
            doc.add_heading(payload, level=1)
        elif kind == "h2":
            doc.add_heading(payload, level=2)
        elif kind == "h3":
            doc.add_heading(payload, level=3)
        elif kind == "h4":
            doc.add_heading(payload, level=4)
        elif kind == "para":
            if payload.startswith("面向《人工智能创新开发实践》"):
                p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                r = p.add_run(payload); r.bold = True; r.font.size = Pt(12)
                continue
            p = doc.add_paragraph(); add_runs(p, payload)
        elif kind == "bullet":
            p = doc.add_paragraph(style="List Bullet"); add_runs(p, payload)
            p.paragraph_format.space_after = Pt(2)
        elif kind == "check":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.55)
            p.paragraph_format.first_line_indent = Cm(-0.55)
            add_runs(p, "\u2610  " + payload)
        elif kind == "num":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.65)
            p.paragraph_format.first_line_indent = Cm(-0.65)
            add_runs(p, payload[0] + ".  " + payload[1])
        elif kind == "table":
            headers, rows = payload
            add_table(doc, headers, rows)
        elif kind == "code":
            for line in payload:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0
                p.paragraph_format.left_indent = Cm(0.3)
                r = p.add_run(line if line else " ")
                r.font.name = "Consolas"
                r._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
                r.font.size = Pt(8.5)
                shade_paragraph(p, "F4F6F7")
            doc.add_paragraph().paragraph_format.space_after = Pt(0)

        if kind == "h2" and payload.startswith("7 模块设计"):
            doc.add_picture(str(flow), width=Cm(16.4))
            cap = doc.add_paragraph(); cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = cap.add_run("图 1 RAG 模块离线索引与在线检索数据流")
            r.bold = True; r.font.size = Pt(9)
            cap.paragraph_format.space_after = Pt(6)

    doc.save(OUT)
    print("saved", OUT)
    print("blocks", len(blocks))


if __name__ == "__main__":
    build()
